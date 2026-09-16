#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Read-only extract of GOJ competency PDF + Factorial assignment importer."""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
OUT_DIR = ROOT / "clients" / "statistic-institute" / "run_log"
OUT_DIR.mkdir(parents=True, exist_ok=True)

XLSX = Path(r"c:\Users\victo\Downloads\competencies_competency_assignment_importer (1).xlsx")
PDF = Path(
    r"c:\Users\victo\Downloads\GOJ-Competency-Model-Competency-Framework-Professional-Pathways-Website-Development.pdf"
)


def ensure_pkgs() -> None:
    for pkg in ("openpyxl", "pypdf"):
        try:
            __import__(pkg)
        except ImportError:
            import subprocess

            subprocess.check_call([sys.executable, "-m", "pip", "install", pkg, "-q"])


def extract_excel() -> str:
    import openpyxl

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    lines = [f"FILE: {XLSX.name}", f"SHEETS: {wb.sheetnames}", ""]
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        lines.append(f"=== SHEET: {name} rows={len(rows)} ===")
        for i, row in enumerate(rows[:80], 1):
            vals = ["" if v is None else str(v).strip() for v in row]
            while vals and vals[-1] == "":
                vals.pop()
            if any(vals):
                lines.append(f"{i}: " + " | ".join(vals))
        if len(rows) > 80:
            lines.append(f"... ({len(rows) - 80} more rows)")
        lines.append("")
    return "\n".join(lines)


def extract_pdf() -> str:
    import pypdf

    reader = pypdf.PdfReader(str(PDF))
    parts = [f"FILE: {PDF.name}", f"PAGES: {len(reader.pages)}", ""]
    for i, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        parts.append(f"----- PAGE {i} -----")
        parts.append(text)
        parts.append("")
    return "\n".join(parts)


def main() -> int:
    ensure_pkgs()
    xlsx_text = extract_excel()
    pdf_text = extract_pdf()
    xlsx_out = OUT_DIR / "_extract_competency_assignment_importer.txt"
    pdf_out = OUT_DIR / "_extract_GOJ_Competency_Framework.txt"
    xlsx_out.write_text(xlsx_text, encoding="utf-8")
    pdf_out.write_text(pdf_text, encoding="utf-8", errors="replace")
    print(xlsx_text[:8000])
    print("\n===== PDF HEAD =====\n")
    print(pdf_text[:10000])
    print(f"\nWROTE {xlsx_out}")
    print(f"WROTE {pdf_out} chars={len(pdf_text)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
